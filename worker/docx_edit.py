"""Format-preserving edits to an existing Word document.

Every paragraph (body, table cells, headers and footers) gets a stable id from its position, so
the editing model can address exactly the text it means. Edits change runs, never styles: a
replacement inherits the formatting of the run where it starts, a rewritten paragraph keeps its
paragraph properties and the first run's character formatting, and inserted paragraphs or rows
are clones of their neighbours (numbering, indents, borders and shading come along).

Operations
  {"type": "replace_text", "find": "...", "replace": "...", "id": "p12"?, "all": bool?}
  {"type": "set_text", "id": "p12", "text": "... **bold** *italic* ..."}
  {"type": "insert_after" | "insert_before", "id": "p12", "text": "...", "like": "p9"?}
  {"type": "delete", "id": "p12"}
  {"type": "fill_blank", "id": "p4"?, "label": "Name:"?, "value": "..."}
  {"type": "check", "id": "p7"?, "label": "I agree"?, "checked": true}
  {"type": "set_cell", "table": 1, "row": 2, "col": 3, "text": "..."}       (1-based)
  {"type": "add_row", "table": 1, "after": 3, "cells": ["...", "..."]}
  {"type": "delete_row", "table": 1, "row": 4}
  {"type": "fill_control", "tag": "...", "value": "..."}                     (content controls)
"""

import copy
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

BLANK = re.compile(r"_{3,}|\.{5,}|…{2,}|…{2,}")
BOXES = {"☐": "☒", "□": "☒", "☒": "☐", "☑": "☐", "■": "□"}


class DocxEditError(RuntimeError):
    pass


def _iter_block(parent, prefix, out, tables):
    """Paragraphs in document order with ids; tables numbered across the document."""
    element = parent.element.body if hasattr(parent, "element") and hasattr(parent.element, "body") else parent._element
    count = 0
    for child in element.iterchildren():
        if child.tag == qn("w:p"):
            count += 1
            out.append((f"{prefix}p{count}", Paragraph(child, parent), None))
        elif child.tag == qn("w:tbl"):
            table = Table(child, parent)
            tables.append(table)
            number = len(tables)
            for r, row in enumerate(table.rows, start=1):
                seen = set()
                for c, cell in enumerate(row.cells, start=1):
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p, paragraph in enumerate(cell.paragraphs, start=1):
                        out.append((f"t{number}.r{r}.c{c}.p{p}", paragraph, (number, r, c)))
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is None:
                continue
            for node in content.iterchildren():
                if node.tag == qn("w:p"):
                    count += 1
                    out.append((f"{prefix}p{count}", Paragraph(node, parent), None))


def paragraphs(document):
    out = []
    tables = []
    _iter_block(document, "", out, tables)
    for index, section in enumerate(document.sections, start=1):
        for kind, part in (("h", section.header), ("f", section.footer)):
            try:
                if part.is_linked_to_previous and index > 1:
                    continue
                for p, paragraph in enumerate(part.paragraphs, start=1):
                    out.append((f"{kind}{index}.p{p}", paragraph, None))
            except Exception:
                continue
    return out, tables


def numbered(paragraph):
    """True when Word numbers or bullets the paragraph itself (numPr on it or on its style chain)."""
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.find(qn("w:numPr")) is not None:
        return True
    style = paragraph.style
    for _ in range(10):
        if style is None:
            return False
        s_pr = style.element.find(qn("w:pPr"))
        if s_pr is not None and s_pr.find(qn("w:numPr")) is not None:
            return True
        style = style.base_style
    return False


def outline(path, max_chars=120_000):
    document = Document(str(path))
    entries, tables = paragraphs(document)
    lines = []
    total = 0
    for pid, paragraph, cell in entries:
        text = paragraph.text
        if not text.strip() and cell is None:
            continue
        entry = {"id": pid, "text": text}
        style = paragraph.style.name if paragraph.style is not None else ""
        if style and style != "Normal":
            entry["style"] = style
        if BLANK.search(text):
            entry["blank"] = True
        if any(box in text for box in BOXES):
            entry["checkbox"] = True
        if numbered(paragraph):
            # Word draws the list number ("1.", "a)"); it is not part of the paragraph's text.
            entry["list"] = True
        total += len(text) + 30
        if total > max_chars:
            break
        lines.append(entry)
    controls = []
    for sdt in document.element.body.iter(qn("w:sdt")):
        props = sdt.find(qn("w:sdtPr"))
        if props is None:
            continue
        tag = props.find(qn("w:tag"))
        alias = props.find(qn("w:alias"))
        name = (tag.get(qn("w:val")) if tag is not None else "") or (alias.get(qn("w:val")) if alias is not None else "")
        if name:
            text = "".join(node.text or "" for node in sdt.iter(qn("w:t")))
            controls.append({"tag": name, "value": text[:200]})
    return {
        "kind": "docx",
        "paragraphs": lines,
        "tables": [{"table": i + 1, "rows": len(t.rows), "cols": len(t.columns)} for i, t in enumerate(tables)],
        "controls": controls[:60],
        "has_text": bool(lines),
    }


# ------------------------------------------------------------------------------------------------
# Run-level helpers


def replace_in_paragraph(paragraph, find, replace, replace_all=True):
    """Replace inside a paragraph; a match spanning runs is written into the run where it starts."""
    runs = list(paragraph.runs)
    text = "".join(run.text for run in runs)
    if not find or find not in text:
        return 0
    starts = {}
    covered = set()
    position = text.find(find)
    while position >= 0:
        starts[position] = True
        covered.update(range(position, position + len(find)))
        if not replace_all:
            break
        position = text.find(find, position + len(find))
    offset = 0
    for run in runs:
        pieces = []
        for index in range(offset, offset + len(run.text)):
            if index in starts:
                pieces.append(replace)
            elif index not in covered:
                pieces.append(text[index])
        offset += len(run.text)
        new_text = "".join(pieces)
        if new_text != run.text:
            run.text = new_text
    return len(starts)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def set_paragraph_text(paragraph, text):
    """Rewrite a paragraph: keeps paragraph properties and the first run's formatting."""
    runs = list(paragraph.runs)
    base = copy.deepcopy(runs[0]._r.find(qn("w:rPr"))) if runs and runs[0]._r.find(qn("w:rPr")) is not None else None
    # Remove runs and hyperlinks, keep bookmarks and paragraph properties.
    for child in list(paragraph._p):
        if child.tag in (qn("w:r"), qn("w:hyperlink"), qn("w:ins"), qn("w:del"), qn("w:smartTag")):
            paragraph._p.remove(child)
    for part in INLINE.split(str(text)):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = not bold and part.startswith("*") and part.endswith("*") and len(part) > 2
        value = part[2:-2] if bold else part[1:-1] if italic else part
        run = paragraph.add_run(value)
        if base is not None:
            existing = run._r.find(qn("w:rPr"))
            if existing is not None:
                run._r.remove(existing)
            run._r.insert(0, copy.deepcopy(base))
        if bold:
            run.bold = True
        if italic:
            run.italic = True


def clone_paragraph(paragraph, text, before=False):
    new = copy.deepcopy(paragraph._p)
    if before:
        paragraph._p.addprevious(new)
    else:
        paragraph._p.addnext(new)
    clone = Paragraph(new, paragraph._parent)
    set_paragraph_text(clone, text)
    return clone


def fill_blank(paragraph, value, label=None):
    text = paragraph.text
    start = 0
    if label:
        position = text.lower().find(label.lower())
        if position < 0:
            return 0
        start = position + len(label)
    match = BLANK.search(text, start)
    if match:
        return replace_in_paragraph(paragraph, match.group(0), value, replace_all=False)
    if label:
        # "Name:" with nothing after it: the value follows the label.
        if not text[start:].strip():
            runs = paragraph.runs
            if runs:
                runs[-1].text = runs[-1].text.rstrip() + " " + value
                return 1
    return 0


def toggle_box(paragraph, checked):
    runs = paragraph.runs
    for run in runs:
        for box in ("☐", "□", "☒", "☑", "■"):
            if box in run.text:
                want = "☒" if checked else "☐"
                run.text = run.text.replace(box, want, 1)
                return 1
    # Word's own checkbox content controls.
    for node in paragraph._p.iter("{http://schemas.microsoft.com/office/word/2010/wordml}checked"):
        node.set("{http://schemas.microsoft.com/office/word/2010/wordml}val", "1" if checked else "0")
        for t in paragraph._p.iter(qn("w:t")):
            if t.text in ("☐", "☒"):
                t.text = "☒" if checked else "☐"
        return 1
    return 0


PART_ROOTS = {qn("w:document"), qn("w:hdr"), qn("w:ftr")}


def _attached(element):
    """Whether an element is still in its part (not removed by an earlier operation)."""
    root = element
    while root.getparent() is not None:
        root = root.getparent()
    return root.tag in PART_ROOTS


def apply_operations(source, output, operations):
    document = Document(str(source))
    warnings = []
    applied = 0

    # Ids name paragraphs, tables and rows as the outline showed them. They are resolved against
    # that original numbering for the whole batch: an insertion or deletion earlier in the batch
    # never shifts what a later operation's id points at.
    entries, tables = paragraphs(document)
    by_id = {pid: paragraph for pid, paragraph, _ in entries}
    table_rows = [[row._tr for row in table.rows] for table in tables]

    def paragraph_for(pid):
        paragraph = by_id.get(pid)
        if paragraph is None:
            raise DocxEditError(f"paragraph not found: {pid}")
        if not _attached(paragraph._p):
            raise DocxEditError(f"paragraph {pid} was removed earlier in this edit")
        return paragraph

    def live_paragraphs():
        return [paragraph for _, paragraph, _ in paragraphs(document)[0]]

    def row_for(number, row):
        if row < 1 or row > len(table_rows[number - 1]):
            raise DocxEditError(f"row not found: table {number} row {row}")
        tr = table_rows[number - 1][row - 1]
        if not _attached(tr):
            raise DocxEditError(f"table {number} row {row} was removed earlier in this edit")
        return tr

    for op in operations[:200]:
        if not isinstance(op, dict):
            continue
        kind = str(op.get("type") or op.get("op") or "")
        pid = str(op.get("id") or "")
        try:
            if kind == "replace_text":
                find = str(op.get("find") or "")
                replace = str(op.get("replace") or "")
                # An explicit id limits the edit to that paragraph; an unknown id is an error,
                # never a licence to search the whole file.
                targets = [paragraph_for(pid)] if pid else live_paragraphs()
                count = 0
                for paragraph in targets:
                    count += replace_in_paragraph(paragraph, find, replace, op.get("all", True) is not False)
                    if count and op.get("all") is False:
                        break
                if not count:
                    warnings.append(f"text not found: {find[:60]}")
                applied += 1 if count else 0
            elif kind == "set_text":
                set_paragraph_text(paragraph_for(pid), str(op.get("text") or ""))
                applied += 1
            elif kind in ("insert_after", "insert_before"):
                anchor = paragraph_for(pid)
                like = str(op.get("like") or "")
                model = paragraph_for(like) if like else anchor
                texts = op.get("texts") if isinstance(op.get("texts"), list) else [op.get("text") or ""]
                current = anchor
                for value in texts[:40]:
                    new = copy.deepcopy(model._p)
                    if kind == "insert_before":
                        anchor._p.addprevious(new)
                    else:
                        current._p.addnext(new)
                    clone = Paragraph(new, anchor._parent)
                    set_paragraph_text(clone, str(value))
                    current = clone
                applied += 1
            elif kind == "delete":
                element = paragraph_for(pid)._p
                element.getparent().remove(element)
                applied += 1
            elif kind == "fill_blank":
                value = str(op.get("value") or "")
                label = op.get("label")
                targets = [paragraph_for(pid)] if pid else [paragraph for paragraph in live_paragraphs() if not label or str(label).lower() in paragraph.text.lower()]
                done = 0
                for paragraph in targets:
                    done = fill_blank(paragraph, value, label)
                    if done:
                        break
                if not done:
                    warnings.append(f"blank not found: {label or pid}")
                applied += done
            elif kind == "check":
                label = op.get("label")
                targets = [paragraph_for(pid)] if pid else [paragraph for paragraph in live_paragraphs() if label and str(label).lower() in paragraph.text.lower()]
                done = 0
                for paragraph in targets:
                    done = toggle_box(paragraph, op.get("checked", True) is not False)
                    if done:
                        break
                if not done:
                    warnings.append(f"checkbox not found: {label or pid}")
                applied += done
            elif kind in ("set_cell", "add_row", "delete_row"):
                number = int(op.get("table") or 0)
                if number < 1 or number > len(tables):
                    raise DocxEditError(f"table not found: {number}")
                table = tables[number - 1]
                if kind == "set_cell":
                    tr = row_for(number, int(op.get("row") or 0))
                    col = int(op.get("col") or 0)
                    current_rows = [row._tr for row in table.rows]
                    cells = table.rows[current_rows.index(tr)].cells
                    # table.cell() indexes a flat grid, so an out-of-range column would land in the next row.
                    if col < 1 or col > len(cells):
                        raise DocxEditError(f"column not found: table {number}, row {op.get('row')}, column {col}")
                    cell = cells[col - 1]
                    first = cell.paragraphs[0]
                    for extra in cell.paragraphs[1:]:
                        extra._p.getparent().remove(extra._p)
                    set_paragraph_text(first, str(op.get("text") or ""))
                elif kind == "add_row":
                    rows = table_rows[number - 1]
                    after = int(op.get("after") or len(rows))
                    model = row_for(number, min(max(after, 1), len(rows)))
                    new = copy.deepcopy(model)
                    model.addnext(new)
                    cells = op.get("cells") if isinstance(op.get("cells"), list) else []
                    for index_, tc in enumerate(new.iterchildren(qn("w:tc"))):
                        ps = list(tc.iterchildren(qn("w:p")))
                        for extra in ps[1:]:
                            tc.remove(extra)
                        if ps:
                            set_paragraph_text(Paragraph(ps[0], table), str(cells[index_]) if index_ < len(cells) else "")
                else:
                    tr = row_for(number, int(op.get("row") or 0))
                    tr.getparent().remove(tr)
                applied += 1
            elif kind == "fill_control":
                name = str(op.get("tag") or op.get("name") or "")
                value = str(op.get("value") or "")
                done = 0
                for sdt in document.element.body.iter(qn("w:sdt")):
                    props = sdt.find(qn("w:sdtPr"))
                    if props is None:
                        continue
                    names = [node.get(qn("w:val")) for node in (props.find(qn("w:tag")), props.find(qn("w:alias"))) if node is not None]
                    if name not in names:
                        continue
                    texts = list(sdt.iter(qn("w:t")))
                    if texts:
                        texts[0].text = value
                        for extra in texts[1:]:
                            extra.text = ""
                    showing = props.find(qn("w:showingPlcHdr"))
                    if showing is not None:
                        props.remove(showing)
                    done = 1
                    break
                if not done:
                    warnings.append(f"content control not found: {name}")
                applied += done
            else:
                warnings.append(f"unknown operation: {kind}")
        except (DocxEditError, IndexError, ValueError) as exc:
            warnings.append(str(exc)[:200])
    if not applied:
        raise DocxEditError("no_edits_applied: " + "; ".join(warnings[:4]))
    document.save(str(output))
    return applied, warnings
