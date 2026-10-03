import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Pt

from worker import docx_edit, pdf_edit


def make_docx(path):
    doc = Document()
    doc.add_heading("Library Card Application", 1)
    para = doc.add_paragraph()
    label = para.add_run("Full name: ")
    label.bold = True
    blank = para.add_run("____________________")
    blank.font.size = Pt(12)
    doc.add_paragraph("Student ID: ____________")
    doc.add_paragraph("The fee is 10 dollars.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Department"
    table.cell(1, 0).text = "Year"
    doc.add_paragraph("☐ I agree to the library rules.")
    doc.save(path)


class DocxEditTest(unittest.TestCase):
    def test_outline_gives_stable_ids_blanks_and_checkboxes(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.docx"
            make_docx(source)
            outline = docx_edit.outline(source)
            by_id = {entry["id"]: entry for entry in outline["paragraphs"]}
            self.assertTrue(by_id["p2"]["blank"])
            self.assertEqual(by_id["p1"]["style"], "Heading 1")
            self.assertTrue(any(entry.get("checkbox") for entry in outline["paragraphs"]))
            self.assertIn("t1.r1.c1.p1", by_id)

    def test_edits_keep_run_formatting(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.docx"
            output = Path(tmp) / "out.docx"
            make_docx(source)
            applied, warnings = docx_edit.apply_operations(source, output, [
                {"type": "fill_blank", "id": "p2", "value": "Arjun Mehta"},
                {"type": "fill_blank", "label": "Student ID:", "value": "20231187"},
                {"type": "replace_text", "find": "10 dollars", "replace": "12 dollars"},
                {"type": "set_cell", "table": 1, "row": 1, "col": 2, "text": "Computer Science"},
                {"type": "check", "label": "I agree", "checked": True},
                {"type": "insert_after", "id": "p4", "text": "A **new** line."},
                {"type": "replace_text", "find": "not in the document", "replace": "x"},
            ])
            self.assertEqual(applied, 6)
            self.assertEqual(len(warnings), 1)
            doc = Document(output)
            name = doc.paragraphs[1]
            self.assertEqual(name.text, "Full name: Arjun Mehta")
            self.assertTrue(name.runs[0].bold)
            self.assertEqual(name.runs[1].font.size, Pt(12))
            texts = [p.text for p in doc.paragraphs]
            self.assertIn("Student ID: 20231187", texts)
            self.assertIn("The fee is 12 dollars.", texts)
            self.assertIn("A new line.", texts)
            self.assertIn("☒ I agree to the library rules.", texts)
            self.assertEqual(doc.tables[0].cell(0, 1).text, "Computer Science")

    def test_no_applicable_operation_is_an_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.docx"
            make_docx(source)
            with self.assertRaises(docx_edit.DocxEditError):
                docx_edit.apply_operations(source, Path(tmp) / "out.docx", [{"type": "delete", "id": "p99"}])

    def test_ids_stay_on_their_paragraphs_through_a_batch(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "abc.docx"
            output = Path(tmp) / "out.docx"
            doc = Document()
            for text in ("Alpha", "Beta", "Gamma"):
                doc.add_paragraph(text)
            doc.save(source)
            docx_edit.apply_operations(source, output, [
                {"type": "delete", "id": "p1"},
                {"type": "set_text", "id": "p2", "text": "New Beta"},
                {"type": "insert_after", "id": "p2", "text": "After Beta"},
                {"type": "replace_text", "id": "p3", "find": "Gamma", "replace": "Delta"},
            ])
            self.assertEqual([p.text for p in Document(output).paragraphs], ["New Beta", "After Beta", "Delta"])

    def test_unknown_or_removed_ids_never_widen_the_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.docx"
            output = Path(tmp) / "out.docx"
            make_docx(source)
            applied, warnings = docx_edit.apply_operations(source, output, [
                {"type": "replace_text", "id": "p999", "find": "10 dollars", "replace": "99 dollars"},
                {"type": "fill_blank", "id": "p998", "value": "X"},
                {"type": "check", "id": "p997", "checked": True},
                {"type": "delete", "id": "p4"},
                {"type": "set_text", "id": "p4", "text": "should not land anywhere"},
                {"type": "set_cell", "table": 1, "row": 2, "col": 2, "text": "Second"},
            ])
            self.assertEqual(applied, 2)
            self.assertEqual(len(warnings), 4)
            texts = [p.text for p in Document(output).paragraphs]
            self.assertNotIn("The fee is 99 dollars.", texts)
            self.assertNotIn("The fee is 10 dollars.", texts)
            self.assertNotIn("should not land anywhere", texts)
            self.assertIn("☐ I agree to the library rules.", texts)
            self.assertIn("Student ID: ____________", texts)

    def test_table_rows_keep_their_numbers_through_a_batch(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.docx"
            output = Path(tmp) / "out.docx"
            make_docx(source)
            docx_edit.apply_operations(source, output, [
                {"type": "add_row", "table": 1, "after": 1, "cells": ["Inserted", ""]},
                {"type": "set_cell", "table": 1, "row": 2, "col": 2, "text": "Third"},
            ])
            table = Document(output).tables[0]
            self.assertEqual([row.cells[0].text for row in table.rows], ["Department", "Inserted", "Year"])
            self.assertEqual(table.cell(2, 1).text, "Third")

    def test_out_of_range_column_never_lands_in_another_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.docx"
            output = Path(tmp) / "out.docx"
            make_docx(source)
            applied, warnings = docx_edit.apply_operations(source, output, [
                {"type": "set_cell", "table": 1, "row": 1, "col": 3, "text": "Wrong"},
                {"type": "set_cell", "table": 1, "row": 1, "col": 0, "text": "Wrong"},
                {"type": "set_cell", "table": 1, "row": 1, "col": 2, "text": "Right"},
            ])
            self.assertEqual(applied, 1)
            self.assertEqual(len(warnings), 2)
            table = Document(output).tables[0]
            self.assertEqual([[cell.text for cell in row.cells] for row in table.rows], [["Department", "Right"], ["Year", ""]])

    def test_outline_marks_paragraphs_word_numbers(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "list.docx"
            doc = Document()
            doc.add_paragraph("Intro.")
            doc.add_paragraph("First point.", style="List Number")
            doc.add_paragraph("2. Typed by hand.")
            doc.save(source)
            by_text = {entry["text"]: entry for entry in docx_edit.outline(source)["paragraphs"]}
            self.assertTrue(by_text["First point."].get("list"))
            self.assertFalse(by_text["2. Typed by hand."].get("list"))
            self.assertFalse(by_text["Intro."].get("list"))


def make_form_pdf(path):
    # Two pages, each with a text field named "Name" under a different parent: student.Name, guardian.Name.
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

    writer = PdfWriter()
    fields = ArrayObject()
    for parent_name in ("student", "guardian"):
        page = writer.add_blank_page(595, 842)
        parent = DictionaryObject({NameObject("/T"): TextStringObject(parent_name), NameObject("/Kids"): ArrayObject()})
        parent_ref = writer._add_object(parent)
        widget = DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject("Name"),
            NameObject("/Rect"): ArrayObject([FloatObject(100), FloatObject(700), FloatObject(300), FloatObject(720)]),
            NameObject("/Parent"): parent_ref,
        })
        widget_ref = writer._add_object(widget)
        parent[NameObject("/Kids")].append(widget_ref)
        page[NameObject("/Annots")] = ArrayObject([widget_ref])
        fields.append(parent_ref)
    writer._root_object[NameObject("/AcroForm")] = writer._add_object(DictionaryObject({NameObject("/Fields"): fields}))
    with open(path, "wb") as handle:
        writer.write(handle)


def make_shared_field_pdf(path, pages=2, rect=True):
    # One text field "Name" (value OLD) with a widget on each page, as forms that repeat a name do.
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

    writer = PdfWriter()
    field = DictionaryObject({
        NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject("Name"),
        NameObject("/V"): TextStringObject("OLD"), NameObject("/Kids"): ArrayObject(),
    })
    field_ref = writer._add_object(field)
    for index in range(pages):
        page = writer.add_blank_page(595, 842)
        widget = DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/Parent"): field_ref,
        })
        if rect or index == 0:
            widget[NameObject("/Rect")] = ArrayObject([FloatObject(100), FloatObject(700), FloatObject(300), FloatObject(720)])
        widget_ref = writer._add_object(widget)
        field[NameObject("/Kids")].append(widget_ref)
        page[NameObject("/Annots")] = ArrayObject([widget_ref])
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica"),
    }))
    writer._root_object[NameObject("/AcroForm")] = writer._add_object(DictionaryObject({
        NameObject("/Fields"): ArrayObject([field_ref]), NameObject("/DA"): TextStringObject("/Helv 0 Tf 0 g"),
        NameObject("/DR"): DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/Helv"): font})}),
    }))
    with open(path, "wb") as handle:
        writer.write(handle)


def make_pdf(path, rows=((780, "Field Trip Permission Form"), (740, "Student name: __________"), (720, "Grade: 7"))):
    # A small text PDF drawn with PDFium itself, so the test needs no other tool.
    import ctypes
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    pdf = pdfium.PdfDocument.new()
    page = pdf.new_page(595, 842)
    font = raw.FPDFText_LoadStandardFont(pdf.raw, b"Helvetica")
    for y, text in rows:
        obj = raw.FPDFPageObj_CreateTextObj(pdf.raw, font, ctypes.c_float(12))
        wide, keep = pdf_edit._wide(text)
        raw.FPDFText_SetText(obj, wide)
        matrix = raw.FS_MATRIX(1, 0, 0, 1, 50, y)
        raw.FPDFPageObj_SetMatrix(obj, ctypes.byref(matrix))
        raw.FPDFPage_InsertObject(page.raw, obj)
    raw.FPDFPage_GenerateContent(page.raw)
    pdf.save(str(path))
    pdf.close()


class PdfEditTest(unittest.TestCase):
    def test_outline_lists_lines_with_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.pdf"
            make_pdf(source)
            outline = pdf_edit.outline(source)
            texts = [line["text"] for line in outline["pages"][0]["lines"]]
            self.assertIn("Student name: __________", texts)
            self.assertTrue(outline["has_text"])

    def test_replace_and_insert_keep_other_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.pdf"
            output = Path(tmp) / "out.pdf"
            make_pdf(source)
            applied, _warnings = pdf_edit.apply_operations(source, output, [
                {"type": "replace_text", "find": "Student name: __________", "replace": "Student name: Jane Doe"},
                {"type": "insert_text", "page": 1, "anchor": "Grade:", "text": "B"},
            ])
            self.assertEqual(applied, 2)
            texts = [line["text"] for line in pdf_edit.outline(output)["pages"][0]["lines"]]
            self.assertIn("Student name: Jane Doe", texts)
            self.assertIn("Field Trip Permission Form", texts)
            self.assertTrue(any("Grade: 7" in text for text in texts))
            self.assertTrue(any(text.strip() == "B" or text.endswith("B") for text in texts))

    def test_line_ids_stay_put_after_an_earlier_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.pdf"
            output = Path(tmp) / "out.pdf"
            make_pdf(source)
            ids = {line["text"]: line["id"] for line in pdf_edit.outline(source)["pages"][0]["lines"]}
            applied, _warnings = pdf_edit.apply_operations(source, output, [
                {"type": "replace_text", "line": ids["Student name: __________"], "find": "__________", "replace": "Jane Doe"},
                {"type": "replace_text", "line": ids["Grade: 7"], "find": "7", "replace": "8"},
                {"type": "rewrite_lines", "lines": [ids["Field Trip Permission Form"]], "text": "Museum Trip Permission Form"},
            ])
            self.assertEqual(applied, 3)
            texts = [line["text"] for line in pdf_edit.outline(output)["pages"][0]["lines"]]
            self.assertIn("Student name: Jane Doe", texts)
            self.assertIn("Grade: 8", texts)
            self.assertIn("Museum Trip Permission Form", texts)

    def test_unknown_line_id_is_not_a_page_wide_edit(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.pdf"
            make_pdf(source)
            with self.assertRaises(pdf_edit.PdfEditError):
                pdf_edit.apply_operations(source, Path(tmp) / "out.pdf", [{"type": "replace_text", "line": "p1.l99", "find": "7", "replace": "8"}])

    def test_anchor_insert_lands_on_the_given_line(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "names.pdf"
            output = Path(tmp) / "out.pdf"
            make_pdf(source, rows=((780, "Name: First"), (700, "Name: Second")))
            lines = pdf_edit.outline(source)["pages"][0]["lines"]
            ids = {line["text"]: line for line in lines}
            self.assertIn("w", ids["Name: Second"])
            pdf_edit.apply_operations(source, output, [
                {"type": "insert_text", "page": 1, "anchor": "Name:", "line": ids["Name: Second"]["id"], "text": "MARK"},
            ])
            after = pdf_edit.outline(output)["pages"][0]["lines"]
            mark = next(line for line in after if "MARK" in line["text"])
            self.assertAlmostEqual(mark["y"], ids["Name: Second"]["y"], delta=3)
            with self.assertRaises(pdf_edit.PdfEditError):
                pdf_edit.apply_operations(source, Path(tmp) / "bad.pdf", [
                    {"type": "insert_text", "page": 1, "anchor": "Grade:", "line": ids["Name: Second"]["id"], "text": "X"},
                ])

    def test_fields_report_where_each_widget_is(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.pdf"
            make_form_pdf(source)
            fields = {field["name"]: field for field in pdf_edit.outline(source)["fields"]}
            self.assertEqual(set(fields), {"student.Name", "guardian.Name"})
            self.assertEqual(fields["student.Name"]["widgets"], [{"page": 1, "x": 100.0, "y": 122.0, "w": 200.0, "h": 20.0}])
            self.assertEqual(fields["guardian.Name"]["widgets"][0]["page"], 2)

    def test_a_field_shown_on_several_pages_reports_every_widget(self):
        # A value fills every widget, so the outline lists them all for the selection check.
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "shared.pdf"
            make_shared_field_pdf(source)
            field = pdf_edit.outline(source)["fields"][0]
            self.assertEqual([widget["page"] for widget in field["widgets"]], [1, 2])
            self.assertNotIn("widgets_partial", field)
            output = Path(tmp) / "out.pdf"
            pdf_edit.fill_fields(source, output, {"Name": "CHANGED"})
            from pypdf import PdfReader
            reader = PdfReader(str(output))
            self.assertEqual(reader.get_fields()["Name"].get("/V"), "CHANGED")
            many = Path(tmp) / "many.pdf"
            make_shared_field_pdf(many, pages=pdf_edit.WIDGET_LIMIT + 1)
            self.assertTrue(pdf_edit.outline(many)["fields"][0]["widgets_partial"])
            unplaced = Path(tmp) / "unplaced.pdf"
            make_shared_field_pdf(unplaced, rect=False)
            field = pdf_edit.outline(unplaced)["fields"][0]
            self.assertEqual(len(field["widgets"]), 1)
            self.assertTrue(field["widgets_partial"])

    def test_widgets_and_lines_share_coordinates_on_a_rotated_page(self):
        from pypdf import PdfReader, PdfWriter
        from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

        with tempfile.TemporaryDirectory() as tmp:
            text_pdf = Path(tmp) / "text.pdf"
            source = Path(tmp) / "rotated.pdf"
            make_pdf(text_pdf, rows=((700, "Student name: First"), (500, "Guardian name: Second")))
            writer = PdfWriter(clone_from=PdfReader(str(text_pdf)))
            page = writer.pages[0]
            page.rotate(90)
            fields = ArrayObject()
            for parent_name, y in (("student", 700), ("guardian", 500)):
                parent = DictionaryObject({NameObject("/T"): TextStringObject(parent_name), NameObject("/Kids"): ArrayObject()})
                parent_ref = writer._add_object(parent)
                widget_ref = writer._add_object(DictionaryObject({
                    NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
                    NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject("Name"),
                    NameObject("/Rect"): ArrayObject([FloatObject(200), FloatObject(y - 4), FloatObject(400), FloatObject(y + 14)]),
                    NameObject("/Parent"): parent_ref,
                }))
                parent[NameObject("/Kids")].append(widget_ref)
                page.setdefault(NameObject("/Annots"), ArrayObject()).append(widget_ref)
                fields.append(parent_ref)
            writer._root_object[NameObject("/AcroForm")] = writer._add_object(DictionaryObject({NameObject("/Fields"): fields}))
            with open(source, "wb") as handle:
                writer.write(handle)

            data = pdf_edit.outline(source)
            lines = {line["text"]: line for line in data["pages"][0]["lines"]}
            widgets = {field["name"]: field["widgets"][0] for field in data["fields"]}

            def level(widget, line):
                return widget["y"] <= line["y"] + line["size"] * 1.5 and widget["y"] + widget["h"] >= line["y"] - line["size"] * 0.5

            # The same test the server uses to decide which field sits beside a selected line.
            self.assertTrue(level(widgets["guardian.Name"], lines["Guardian name: Second"]))
            self.assertFalse(level(widgets["student.Name"], lines["Guardian name: Second"]))
            self.assertTrue(level(widgets["student.Name"], lines["Student name: First"]))
            self.assertFalse(level(widgets["guardian.Name"], lines["Student name: First"]))

    def test_missing_text_raises(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "form.pdf"
            make_pdf(source)
            with self.assertRaises(pdf_edit.PdfEditError):
                pdf_edit.apply_operations(source, Path(tmp) / "out.pdf", [{"type": "replace_text", "find": "absent", "replace": "x"}])


if __name__ == "__main__":
    unittest.main()
